"""
Mantiene, dentro de cada carpeta de fotos exportadas, un Excel que dice de
qué PDF/Excel de proveedor salió cada foto (y en qué página/fila) — para
poder rastrear cualquier producto de vuelta a su catálogo original.

Se actualiza solo, agregando filas nuevas cada vez que se exporta un lote
(ver `gui_profesional.py:_exportar`) — nunca rehace el archivo de cero, así
que conserva fotos exportadas en sesiones anteriores.
"""

from __future__ import annotations

import time
from io import BytesIO
from pathlib import Path

import extract_msg
import fitz  # PyMuPDF
import openpyxl
from openpyxl.drawing.image import Image as ImagenExcel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as ImagenPIL

import extraer_excel
import extraer_msg
import extraer_pdf

# Nombre real del archivo/etiqueta interna -> como se quiere ver en la
# columna "Archivo de origen" del índice. Un origen que no está acá se
# muestra tal cual (el nombre real del archivo), para que agregar un
# catálogo nuevo nunca deje una fila en blanco por no estar en esta lista.
ETIQUETAS_ORIGEN = {
    "PACKING LIST #79599.xlsx": "Packing List",
    "BASH CR05-25.xlsx": "Excel BASH",
    "PACKING LIST 2FORTY contenedor 134.xlsx": "Packing List",
    "FOTOS.msg": "Correo Electrónico",
    "Distribuidora Cachos - Osiris FA26.xlsx": "Excel Distribuidora Cachos",
    "Grupo A.xlsx": "Excel Grupo A",
    "imágenes internas de Tu Calzado": "Imágenes internas Tu Calzado",
    "PREVENTA SS2026.pdf": "PDF GALA",
    "(descargada de internet)": "Imágenes de control tomadas de internet",
}


def etiqueta_origen(nombre_archivo: str) -> str:
    return ETIQUETAS_ORIGEN.get(nombre_archivo, nombre_archivo)


NOMBRE_ARCHIVO = "indice_fuente_imagenes.xlsx"
HOJA_DATOS = "Fuente de imágenes"
ENCABEZADOS = ["Foto", "Código de producto", "Archivo de origen", "Ubicación en el origen",
              "Imagen fondo blanco"]
ALTO_FILA_CON_FOTO = 60
ANCHO_MINIATURA = 70


def _codigo_base(stem: str) -> str:
    """`100702BBK_01` (recorte final, con sufijo de posición) -> `100702BBK`
    (el código tal cual aparece en el PDF/Excel de origen).

    Recibe un stem YA sin extensión (`Path(nombre).stem`, aplicado por quien
    llama) — no se le vuelve a aplicar `Path(...).stem` acá, porque estos
    códigos de proveedor a veces tienen puntos de verdad (ej.
    `1122.888.7286-15745`), y `Path` interpreta cualquier punto como
    separador de extensión: cortaba el código en el punto equivocado y la
    búsqueda en el Excel fallaba en silencio (quedaba "no encontrado" pese a
    que el código sí estaba en el catálogo)."""
    partes = stem.rsplit("_", 1)
    if len(partes) == 2 and partes[1].isdigit():
        return partes[0]
    return stem


class _IndiceCatalogos:
    """Busca a qué PDF/Excel pertenece un código, cacheando cada catálogo la
    primera vez que se abre — para no releer los mismos 14 Excel por cada
    foto del lote."""

    def __init__(self, raiz: Path) -> None:
        self.raiz = Path(raiz)
        self._pdf_cache: dict[Path, dict[str, int]] = {}
        self._excel_cache: dict[Path, dict[str, int]] = {}
        self._msg_cache: dict[Path, set[str]] = {}

    def _paginas_pdf(self, ruta: Path) -> dict[str, int]:
        if ruta not in self._pdf_cache:
            mapa: dict[str, int] = {}
            try:
                doc = fitz.open(ruta)
                try:
                    for i in range(len(doc)):
                        cod = extraer_pdf._codigo_de_pagina(doc[i])
                        if cod:
                            mapa.setdefault(cod, i + 1)
                finally:
                    doc.close()
            except Exception:  # noqa: BLE001 - un PDF roto no frena el índice
                pass
            self._pdf_cache[ruta] = mapa
        return self._pdf_cache[ruta]

    def _filas_excel(self, ruta: Path) -> dict[str, int]:
        if ruta not in self._excel_cache:
            mapa: dict[str, int] = {}
            try:
                # SIN read_only: en ese modo cada hoja se puede recorrer una
                # sola vez, así que `_detectar_fila_encabezado` la consumía y
                # el `iter_rows` siguiente ya no devolvía el encabezado — el
                # resultado era 0 códigos y todas las fotos quedaban como
                # "(no encontrado)" en el índice, sin ningún error visible.
                wb = openpyxl.load_workbook(ruta, data_only=True)
                for ws in wb.worksheets:
                    fila_enc = extraer_excel._detectar_fila_encabezado(ws)
                    encabezados = {c.column: extraer_excel._normalizar(c.value)
                                  for c in next(ws.iter_rows(min_row=fila_enc, max_row=fila_enc))
                                  if c.value}
                    col_cod = extraer_excel._elegir_columna_codigo(encabezados)
                    if not col_cod:
                        continue
                    for fila, cod in extraer_excel._mapa_codigos_por_fila(ws, fila_enc, col_cod).items():
                        if cod:
                            mapa.setdefault(cod, fila)
                wb.close()
            except Exception:  # noqa: BLE001 - un Excel roto no frena el índice
                pass
            self._excel_cache[ruta] = mapa
        return self._excel_cache[ruta]

    def _codigos_msg(self, ruta: Path) -> set[str]:
        if ruta not in self._msg_cache:
            codigos: set[str] = set()
            try:
                msg = extract_msg.Message(ruta)
                try:
                    for adjunto in msg.attachments:
                        nombre = adjunto.longFilename or adjunto.shortFilename or ""
                        stem = Path(nombre).stem
                        if stem:
                            codigos.add(extraer_msg._sanitizar_nombre(stem))
                finally:
                    msg.close()
            except Exception:  # noqa: BLE001 - un .msg roto no frena el índice
                pass
            self._msg_cache[ruta] = codigos
        return self._msg_cache[ruta]

    def _buscar_directo(self, codigo: str) -> tuple[str, str] | None:
        for pdf in sorted(self.raiz.rglob("*.pdf")):
            pagina = self._paginas_pdf(pdf).get(codigo)
            if pagina:
                return pdf.name, f"PDF · página {pagina}"
        for xlsx in sorted(self.raiz.rglob("*.xlsx")):
            fila = self._filas_excel(xlsx).get(codigo)
            if fila:
                return xlsx.name, f"Excel · fila {fila}"
        for msg in sorted(self.raiz.rglob("*.msg")):
            if codigo in self._codigos_msg(msg):
                return msg.name, "Correo · adjunto"
        return None

    def buscar(self, codigo: str) -> tuple[str, str] | None:
        if not self.raiz.exists():
            return None
        resultado = self._buscar_directo(codigo)
        if resultado:
            return resultado
        # `extraer_excel.py` le agrega "_N" al código cuando el mismo
        # aparece repetido en el Excel (2da aparición, 3ra, ...) — ese
        # sufijo no es parte del código real, así que si la búsqueda directa
        # falla se prueba sin él (ej. `AL-24701C-BLK_2` -> `AL-24701C-BLK`).
        partes = codigo.rsplit("_", 1)
        if len(partes) == 2 and partes[1].isdigit():
            return self._buscar_directo(partes[0])
        return None


def _miniatura(ruta_foto: Path) -> ImagenExcel | None:
    try:
        im = ImagenPIL.open(ruta_foto)
        if im.mode == "RGBA":
            fondo = ImagenPIL.new("RGB", im.size, (255, 255, 255))
            fondo.paste(im, mask=im.split()[3])
            im = fondo
        else:
            im = im.convert("RGB")
        im.thumbnail((ANCHO_MINIATURA, ANCHO_MINIATURA))
        buf = BytesIO()
        im.save(buf, format="PNG")
        buf.seek(0)
        return ImagenExcel(buf)
    except Exception:  # noqa: BLE001 - una foto ilegible no debe tumbar el índice
        return None


def _guardar_con_reintentos(wb: openpyxl.Workbook, ruta: Path, intentos: int = 5) -> None:
    """`wb.save` puede chocar con un bloqueo de un instante justo después de
    escribir (antivirus escaneando el .xlsx recién creado, u OneDrive/Excel
    soltando el archivo) — pasó de verdad al agregar el gráfico del Resumen:
    el segundo lote de un lote múltiple se perdía en silencio porque la
    excepción no se atrapaba y el resto de la carga nunca corría. Reintenta
    con una pausa corta antes de darse por vencido de verdad."""
    ultimo_error = None
    for intento in range(intentos):
        try:
            wb.save(ruta)
            return
        except PermissionError as exc:
            ultimo_error = exc
            time.sleep(0.4 * (intento + 1))
    raise ultimo_error


def _actualizar_resumen(wb: openpyxl.Workbook, ws_datos) -> None:
    """Hoja "Resumen": cuántas fotos hay de cada fuente, con fórmulas
    `COUNTIF` — no es una tabla dinámica nativa (openpyxl no puede crear
    esas), pero se recalcula sola en Excel cada vez que se abre o se agrega
    una fila nueva a "Fuente de imágenes", que es lo que importa para tener
    control real sin mantenimiento manual."""
    nombre = "Resumen"
    if nombre in wb.sheetnames:
        del wb[nombre]
    wsr = wb.create_sheet(nombre, 0)

    conteo: dict[str, int] = {}
    for r in range(2, ws_datos.max_row + 1):
        v = ws_datos.cell(row=r, column=3).value or "(sin fuente)"
        conteo[v] = conteo.get(v, 0) + 1

    wsr.append(["Fuente", "Cantidad de fotos"])
    for cell in wsr[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5233")
        cell.alignment = Alignment(horizontal="center")

    fila = 2
    for fuente, _n in sorted(conteo.items(), key=lambda x: -x[1]):
        wsr.cell(row=fila, column=1, value=fuente)
        wsr.cell(row=fila, column=2,
                 value=f"=COUNTIF('{ws_datos.title}'!C:C,A{fila})")
        fila += 1

    fila_total = fila
    c1 = wsr.cell(row=fila_total, column=1, value="Total")
    c2 = wsr.cell(row=fila_total, column=2, value=f"=SUM(B2:B{fila_total - 1})")
    c1.font = c2.font = Font(bold=True)

    wsr.column_dimensions["A"].width = 36
    wsr.column_dimensions["B"].width = 18
    wsr.freeze_panes = "A2"


def actualizar_indice(destino: Path, manifiesto_export: dict, raiz_catalogos: Path) -> Path:
    """Agrega al índice de `destino` las fotos de este export que todavía no
    estaban — nunca duplica ni borra filas de exportaciones anteriores."""
    destino = Path(destino)
    ruta_indice = destino / NOMBRE_ARCHIVO

    if ruta_indice.exists():
        wb = openpyxl.load_workbook(ruta_indice)
        # Por NOMBRE, no `wb.active`: `_actualizar_resumen` inserta la hoja
        # "Resumen" en la posición 0, y eso corre "activa" a esa hoja —
        # pasó de verdad: con `wb.active` todo lo que se agregaba después
        # de la primera vez caía en la hoja de Resumen en vez de en la de
        # datos, y `ya_listados` leía basura de ahí (0 códigos reales), así
        # que cada lote nuevo se creía "ya todo listado" y no agregaba nada.
        if HOJA_DATOS in wb.sheetnames:
            ws = wb[HOJA_DATOS]
        else:
            # Un índice de una versión más vieja (o renombrado a mano en
            # Excel) puede no tener la hoja "Fuente de imágenes" con ese
            # nombre exacto. Caer a `wb.active` reintroduce el mismo bug de
            # arriba, porque la hoja activa puede ser "Resumen" — en vez de
            # eso, se toma la primera hoja que NO sea "Resumen" y se le pone
            # el nombre esperado, para no repetir el problema silenciosamente.
            candidatas = [n for n in wb.sheetnames if n != "Resumen"]
            if not candidatas:
                raise ValueError(
                    f"«{ruta_indice.name}» no tiene ninguna hoja de datos reconocible "
                    f"(solo hay: {wb.sheetnames}). Revisalo a mano en Excel.")
            ws = wb[candidatas[0]]
            ws.title = HOJA_DATOS
        ya_listados = {ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)}
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = HOJA_DATOS
        ws.append(ENCABEZADOS)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F5233")
            cell.alignment = Alignment(horizontal="center")
        anchos = [12, 22, 26, 20, 22]
        for i, ancho in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.freeze_panes = "A2"
        ya_listados = set()

    carpetas = manifiesto_export.get("carpetas", {})
    archivos_blanco = carpetas.get("fondo_blanco", {}).get("archivos", [])
    archivos_transp = carpetas.get("fondo_transparente", {}).get("archivos", [])
    codigos = sorted({Path(n).stem for n in (archivos_blanco or archivos_transp)})

    buscador = _IndiceCatalogos(raiz_catalogos)
    agregados = 0
    for codigo in codigos:
        if codigo in ya_listados:
            continue
        origen = buscador.buscar(_codigo_base(codigo))
        archivo_origen, ubicacion = origen if origen else ("(no encontrado)", "")

        fila = ws.max_row + 1
        ws.cell(row=fila, column=2, value=codigo)
        ws.cell(row=fila, column=3, value=etiqueta_origen(archivo_origen))
        ws.cell(row=fila, column=4, value=ubicacion)
        ws.cell(row=fila, column=5, value=f"{codigo}.jpg" if archivos_blanco else "")

        foto_origen = destino / "fondo_blanco" / f"{codigo}.jpg"
        if not foto_origen.exists():
            foto_origen = destino / "fondo_transparente" / f"{codigo}.png"
        if foto_origen.exists():
            img = _miniatura(foto_origen)
            if img is not None:
                ws.row_dimensions[fila].height = ALTO_FILA_CON_FOTO
                ws.add_image(img, f"A{fila}")
        agregados += 1

    if agregados or not ruta_indice.exists():
        _actualizar_resumen(wb, ws)
        _guardar_con_reintentos(wb, ruta_indice)
    return ruta_indice
